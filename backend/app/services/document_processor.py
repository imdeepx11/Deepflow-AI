import os
import logging
import pymupdf as fitz  # PyMuPDF
import docx

logger = logging.getLogger(__name__)

class DocumentProcessor:
    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """
        Extract text content from PDF, DOCX, or TXT files.
        """
        ext = file_type.lower()
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        extracted_text = ""
        try:
            if ext == ".pdf" or file_path.endswith(".pdf"):
                extracted_text = DocumentProcessor._extract_pdf(file_path)
            elif ext in [".docx", ".doc"] or file_path.endswith((".docx", ".doc")):
                extracted_text = DocumentProcessor._extract_docx(file_path)
            elif ext == ".txt" or file_path.endswith(".txt"):
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
            else:
                # Fallback text reading
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
        except Exception as e:
            logger.error(f"Text extraction failed for {file_path}: {e}")
            extracted_text = ""

        return extracted_text.strip()

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        text_parts = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text_parts.append(page.get_text())
        return "\n".join(text_parts)

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        text_parts = []
        for p in doc.paragraphs:
            if p.text:
                text_parts.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        return "\n".join(text_parts)
